from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.web_scrapper_agent.config import LEAD_OUTPUT_DIR


HEADERS = [
    "Lead ID",
    "Business Name",
    "Address",
    "Phone",
    "Email",
    "Website",
    "Sector / Domain",
    "Fit Score",
    "Confidence",
    "AI Approach Comment",
    "Fit Reason",
    "Source URL",
    "Evidence URLs",
    "Status",
    "Scraped At",
]


def export_leads_workbook(
    leads: list[dict[str, Any]],
    profile: dict[str, Any],
    *,
    output_dir: Path = LEAD_OUTPUT_DIR,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    path = output_dir / f"lead_generation_{timestamp}.xlsx"

    workbook = Workbook()
    leads_sheet = workbook.active
    leads_sheet.title = "Leads"
    leads_sheet.append(HEADERS)

    for lead in leads:
        leads_sheet.append(
            [
                lead.get("lead_id", ""),
                lead.get("business_name", ""),
                lead.get("address", ""),
                lead.get("phone", ""),
                lead.get("email", ""),
                lead.get("website", ""),
                lead.get("sector", ""),
                lead.get("fit_score", 0),
                lead.get("confidence", 0),
                lead.get("approach_comment", ""),
                lead.get("fit_reason", ""),
                lead.get("source_url", ""),
                "\n".join(lead.get("evidence_urls", [])),
                lead.get("status", "New"),
                lead.get("scraped_at", ""),
            ]
        )

    _style_leads_sheet(leads_sheet)
    _add_sources_sheet(workbook, leads)
    _add_summary_sheet(workbook, profile, len(leads), timestamp)
    workbook.save(path)
    return path


def _style_leads_sheet(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 30
    widths = {
        "A": 14,
        "B": 28,
        "C": 32,
        "D": 18,
        "E": 30,
        "F": 38,
        "G": 20,
        "H": 12,
        "I": 12,
        "J": 56,
        "K": 48,
        "L": 38,
        "M": 50,
        "N": 14,
        "O": 24,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in ("F", "L"):
            cell = sheet[f"{column}{row[0].row}"]
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = Font(color="0563C1", underline="single")
        sheet[f"I{row[0].row}"].number_format = "0%"

    if sheet.max_row >= 2:
        table = Table(displayName="LeadsTable", ref=f"A1:O{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        sheet.conditional_formatting.add(
            f"H2:H{sheet.max_row}",
            ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
        )

    validation = DataValidation(
        type="list",
        formula1='"New,Contacted,Replied,Qualified,Closed,Not a fit"',
        allow_blank=True,
    )
    sheet.add_data_validation(validation)
    validation.add(f"N2:N{max(sheet.max_row, 2)}")


def _add_sources_sheet(workbook: Workbook, leads: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("Sources")
    sheet.append(["Lead ID", "Business Name", "Evidence URL", "Purpose"])
    for lead in leads:
        for url in lead.get("evidence_urls", []):
            sheet.append(
                [lead.get("lead_id", ""), lead.get("business_name", ""), url, "Public page used for extraction"]
            )
    _style_simple_sheet(sheet, widths={"A": 14, "B": 28, "C": 70, "D": 34})


def _add_summary_sheet(
    workbook: Workbook,
    profile: dict[str, Any],
    lead_count: int,
    timestamp: str,
) -> None:
    sheet = workbook.create_sheet("Run Summary")
    sheet.append(["Field", "Value"])
    rows = [
        ("Profession", profile.get("profession", "")),
        ("Location", profile.get("location", "")),
        ("Services", ", ".join(profile.get("services", []))),
        ("Target sectors", ", ".join(profile.get("target_sectors", []))),
        ("Source URLs", "\n".join(profile.get("source_urls", []))),
        ("Leads generated", lead_count),
        ("Generated at", timestamp),
    ]
    for row in rows:
        sheet.append(row)
    _style_simple_sheet(sheet, widths={"A": 24, "B": 90})


def _style_simple_sheet(sheet: Any, widths: dict[str, int]) -> None:
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
